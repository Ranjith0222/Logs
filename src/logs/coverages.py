from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from logs.ruleset import extract_rulesets

BOP_COVERAGES_CANDIDATES = (
    Path(__file__).resolve().parent / "data" / "BOP_coverages.xlsx",
    Path(__file__).resolve().parents[2] / "templates" / "BOP_coverages.xlsx",
    Path.cwd() / "templates" / "BOP_coverages.xlsx",
)

# Canonical section rulesets that may omit a single BuildingCvgCode.
SECTION_COVERAGE_CODES: dict[str, str] = {
    "Building": "400101",
    "Business Personal Property": "400102",
    "Liability": "400127",
}

# Prefer short section labels in the coverage-wise table.
SECTION_COVERAGE_LABELS: dict[str, str] = {
    "400101": "Building",
    "400102": "Business Personal Property",
    "400127": "Liability",
}

# Liability occurrence limit (400127) resolution order from UW ruleset logic.
LIABILITY_OCCURRENCE_LIMIT_FIELDS: tuple[str, ...] = (
    "FinalBuilding400127SI",
    "Final400127SI",
    "127FinalBCvgSI",
    "T127FinalBCvgSI",
    "ItemCvg400127SI",
    "LiabilityandMedicalExp",
    "LiabilityBCvgSI",
    "400127BUserSI",
    "BUserSI",
)


@dataclass(frozen=True)
class CoverageRow:
    code: str
    label: str
    sum_insured: str | None = None
    base_rate: str | None = None
    final_rate: str | None = None
    premium: str | None = None


def resolve_coverages_path(path: str | Path | None = None) -> Path | None:
    if path:
        candidate = Path(path)
        if not candidate.exists():
            raise FileNotFoundError(f"BOP coverages workbook not found: {candidate}")
        return candidate
    for candidate in BOP_COVERAGES_CANDIDATES:
        if candidate.exists():
            return candidate
    return None


def load_coverage_descriptions(path: str | Path | None = None) -> dict[str, str]:
    """Map coverage code (e.g. 400101) → description from BOP_coverages.xlsx."""
    workbook_path = resolve_coverages_path(path)
    if workbook_path is None:
        return {}
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb.active
    mapping: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or row[0] is None:
            continue
        code = str(row[0]).strip()
        if not code.isdigit():
            continue
        desc = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if desc:
            mapping[code] = desc
    return mapping


def _parse_list_values(raw: str | None) -> list[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text:
        return []
    if text.startswith("[") and text.endswith("]"):
        text = text[1:-1]
    return [part.strip() for part in text.split(",") if part.strip()]


def _single_coverage_code(*candidates: str | None) -> str | None:
    for raw in candidates:
        parts = _parse_list_values(raw)
        if len(parts) == 1 and parts[0].isdigit():
            return parts[0]
        if raw is not None:
            text = str(raw).strip()
            if text.isdigit():
                return text
    return None


def _first_scalar(raw: str | None) -> str | None:
    if raw is None:
        return None
    parts = _parse_list_values(raw)
    if not parts:
        text = str(raw).strip()
        return text or None
    return parts[0]


def is_nonzero_amount(raw: str | None) -> bool:
    """True when a numeric amount is present and not zero."""
    scalar = _first_scalar(raw)
    if scalar is None:
        return False
    try:
        return abs(float(scalar)) > 0
    except ValueError:
        return bool(scalar)


def coverage_row_has_value(row: CoverageRow) -> bool:
    """Keep coverages with a real premium (or core Building/BPP/Liability values).

    Drops rows where Sum Insured / Base / Final / Premium are all zero, and
    drops non-core coverages that only carry default SI/rates with $0 premium.
    """
    has_premium = is_nonzero_amount(row.premium)
    if has_premium:
        return True
    if row.code in SECTION_COVERAGE_CODES.values():
        return (
            is_nonzero_amount(row.sum_insured)
            or is_nonzero_amount(row.base_rate)
            or is_nonzero_amount(row.final_rate)
        )
    return False


def _io_maps(ruleset) -> tuple[dict[str, str], dict[str, str], dict[str, str]]:
    inputs = {item.name: item.value for item in ruleset.inputs}
    outputs = {item.name: item.value for item in ruleset.outputs}
    evaluations = {item.name: item.value for item in ruleset.evaluations}
    return inputs, outputs, evaluations


def resolve_liability_occurrence_limit(
    inputs: dict[str, str],
    outputs: dict[str, str],
    evaluations: dict[str, str] | None = None,
) -> str | None:
    """Liability occurrence limit (400127) from UW FinalBuilding400127SI logic.

    Liability outputs ``BuildingCoverIncrementalSI`` as a comma-separated triple
    (400127 occurrence, 400129 products, 400151 aggregate). Use the dedicated
    400127 final/user SI fields instead of blindly taking the first list value.
    """
    evaluations = evaluations or {}
    for name in LIABILITY_OCCURRENCE_LIMIT_FIELDS:
        for source in (evaluations, inputs, outputs):
            if name not in source:
                continue
            value = _first_scalar(source.get(name))
            if value is not None:
                return value
    # Last resort: first value of the multi-cover SI list (occurrence slot).
    return _first_scalar(
        outputs.get("BuildingCoverIncrementalSI") or outputs.get("CoverageIncrementalSI")
    )


def _merge_si_from_arrays(log_path: str | Path) -> dict[str, str]:
    """Zip parallel BuildingCvgCode / BuildingCoverIncrementalSI arrays when present."""
    extract = extract_rulesets(str(log_path))
    best: dict[str, str] = {}
    for ruleset in extract.rulesets:
        inputs, outputs, _evaluations = _io_maps(ruleset)
        for source in (outputs, inputs):
            codes = _parse_list_values(
                source.get("BuildingCvgCode")
                or source.get("BCvgCode")
                or source.get("CvgCode")
            )
            sis = _parse_list_values(
                source.get("BuildingCoverIncrementalSI")
                or source.get("CoverageIncrementalSI")
            )
            if len(codes) < 2 or len(codes) != len(sis):
                continue
            if not all(code.isdigit() for code in codes):
                continue
            if len(codes) >= len(best):
                best = dict(zip(codes, sis))
    return best


def _coverage_sum_insured(
    code: str,
    ruleset_name: str,
    inputs: dict[str, str],
    outputs: dict[str, str],
    evaluations: dict[str, str],
) -> str | None:
    if code == "400127" or ruleset_name == "Liability":
        return resolve_liability_occurrence_limit(inputs, outputs, evaluations)
    return _first_scalar(
        outputs.get("BuildingCoverIncrementalSI")
        or outputs.get("CoverageIncrementalSI")
        or inputs.get("BUserSI")
        or inputs.get("IUserSI")
    )


def extract_coverage_rows(
    log_path: str | Path,
    *,
    coverages_path: str | Path | None = None,
    include_zero: bool = False,
) -> list[CoverageRow]:
    """Extract coverage-wise Sum Insured / Base Rate / Final Rate / Premium from a log."""
    descriptions = load_coverage_descriptions(coverages_path)
    extract = extract_rulesets(str(log_path))
    by_code: dict[str, CoverageRow] = {}

    for ruleset in extract.rulesets:
        inputs, outputs, evaluations = _io_maps(ruleset)
        code = _single_coverage_code(
            inputs.get("BuildingCvgCode"),
            outputs.get("BuildingCvgCode"),
            inputs.get("CvgCode"),
            outputs.get("CvgCode"),
            inputs.get("BCvgCode"),
            outputs.get("BCvgCode"),
        )
        if code is None:
            code = SECTION_COVERAGE_CODES.get(ruleset.name)
        if code is None:
            continue

        sum_insured = _coverage_sum_insured(
            code, ruleset.name, inputs, outputs, evaluations
        )
        base_rate = _first_scalar(
            outputs.get("BuildingCoverBaseRate")
            or evaluations.get("LiabililtyRate")
            or evaluations.get("BuildingBaseRate")
            or evaluations.get("BPPBaseRate")
        )
        final_rate = _first_scalar(
            outputs.get("BuildingCoverUserRate")
            or evaluations.get("FinalRate")
            or evaluations.get("BuildingFinalRate")
        )
        premium = _first_scalar(
            outputs.get("BuildingCoverIncrementalPrem")
            or outputs.get("CoverageIncrementalPremium")
            or evaluations.get("LiabilityPremium")
        )

        # Skip rulesets that only mention the code with no rating payloads.
        if not any([sum_insured, base_rate, final_rate, premium]):
            continue

        label = (
            SECTION_COVERAGE_LABELS.get(code)
            or descriptions.get(code)
            or ruleset.name
            or code
        )
        existing = by_code.get(code)
        # Prefer later executions that carry rates/premium (rating rulesets).
        score = sum(1 for v in (base_rate, final_rate, premium) if v is not None)
        prev_score = 0
        if existing:
            prev_score = sum(
                1
                for v in (existing.base_rate, existing.final_rate, existing.premium)
                if v is not None
            )
        if existing is None or score >= prev_score:
            by_code[code] = CoverageRow(
                code=code,
                label=label,
                sum_insured=sum_insured or (existing.sum_insured if existing else None),
                base_rate=base_rate or (existing.base_rate if existing else None),
                final_rate=final_rate or (existing.final_rate if existing else None),
                premium=premium or (existing.premium if existing else None),
            )

    # Fill any missing SI from parallel coverage arrays (not for Liability —
    # occurrence limit comes from FinalBuilding400127SI logic above).
    for code, si in _merge_si_from_arrays(log_path).items():
        if code == "400127":
            continue
        if code in by_code:
            if by_code[code].sum_insured is None:
                row = by_code[code]
                by_code[code] = CoverageRow(
                    code=row.code,
                    label=row.label,
                    sum_insured=si,
                    base_rate=row.base_rate,
                    final_rate=row.final_rate,
                    premium=row.premium,
                )
        elif include_zero:
            by_code[code] = CoverageRow(
                code=code,
                label=SECTION_COVERAGE_LABELS.get(code)
                or descriptions.get(code, code),
                sum_insured=si,
            )

    rows = list(by_code.values())
    if not include_zero:
        rows = [row for row in rows if coverage_row_has_value(row)]

    # Stable order: section coverages first, then numeric code order.
    priority = {code: idx for idx, code in enumerate(SECTION_COVERAGE_CODES.values())}

    def sort_key(row: CoverageRow) -> tuple[int, str]:
        return (priority.get(row.code, 100), row.code)

    return sorted(rows, key=sort_key)
