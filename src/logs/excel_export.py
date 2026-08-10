from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookType

from logs.ruleset import pick_field_values
from logs.sections import RATING_SECTION_NAMES, SECTION_SPECS, SectionSpec, all_alias_names

TEMPLATE_CANDIDATES = (
    Path(__file__).resolve().parent / "data" / "policy_rating_template.xlsx",
    Path(__file__).resolve().parents[2] / "templates" / "policy_rating_template.xlsx",
    Path.cwd() / "templates" / "policy_rating_template.xlsx",
)


def resolve_template_path(template: str | Path | None = None) -> Path:
    if template:
        path = Path(template)
        if not path.exists():
            raise FileNotFoundError(f"Excel template not found: {path}")
        return path
    for path in TEMPLATE_CANDIDATES:
        if path.exists():
            return path
    raise FileNotFoundError(
        "Excel template not found. Expected templates/policy_rating_template.xlsx "
        "or logs/data/policy_rating_template.xlsx."
    )


def _to_number(raw: str | None) -> float | int | str | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    try:
        value = float(text)
    except ValueError:
        return text
    if value.is_integer():
        return int(value)
    return value


def _lookup_aliases(values: dict[str, str | None], aliases: tuple[str, ...]) -> str | None:
    for name in aliases:
        if name in values and values[name] is not None and str(values[name]).strip() != "":
            return values[name]
    return None


def collect_section_values(log_path: str | Path, spec: SectionSpec) -> dict[str, str | None]:
    """Extract one rating section and resolve logical factor keys."""
    from logs.ruleset import extract_rulesets

    extract = extract_rulesets(str(log_path), ruleset_name=spec.ruleset_name)
    if not extract.rulesets:
        return {}

    picked = pick_field_values(extract.rulesets[0], all_alias_names(spec))
    raw = {item.name: item.value for item in picked if item.found}

    resolved: dict[str, str | None] = {}
    for key, aliases in spec.field_aliases.items():
        resolved[key] = _lookup_aliases(raw, aliases)
    return resolved


def collect_all_section_values(log_path: str | Path) -> dict[str, dict[str, str | None]]:
    return {spec.ruleset_name: collect_section_values(log_path, spec) for spec in SECTION_SPECS}


def fill_section_factors(
    workbook: WorkbookType,
    spec: SectionSpec,
    values: dict[str, str | None],
) -> None:
    summary = workbook[workbook.sheetnames[0]]
    for key, row in spec.cell_rows.items():
        if key not in values or values[key] is None:
            continue
        cell = f"{spec.column}{row}"
        summary[cell] = _to_number(values[key])


def fill_policy_header(
    workbook: WorkbookType,
    *,
    policy_no: str | None,
    irpm: str | None = None,
) -> None:
    summary = workbook[workbook.sheetnames[0]]
    if policy_no:
        summary["A1"] = policy_no
        try:
            summary.title = policy_no[:31]
        except ValueError:
            pass
    if irpm is not None:
        summary["T2"] = f"With IRPM  {irpm}"


def fill_limits_from_building(
    workbook: WorkbookType,
    log_path: str | Path,
) -> None:
    from logs.ruleset import extract_rulesets

    extract = extract_rulesets(str(log_path), ruleset_name="Building")
    if not extract.rulesets:
        return
    ruleset = extract.rulesets[0]
    inputs = {item.name: item.value for item in ruleset.inputs}
    outputs = {item.name: item.value for item in ruleset.outputs}
    summary = workbook[workbook.sheetnames[0]]
    building_limit = inputs.get("BuildingLimit") or outputs.get("BuildingCoverIncrementalSI")
    bpp_limit = inputs.get("BusinessPersonalPropLimit")
    if building_limit:
        summary["F3"] = _to_number(building_limit)
    if bpp_limit:
        summary["G3"] = _to_number(bpp_limit)
    if outputs.get("BuildingCoverBaseRate"):
        summary["F20"] = _to_number(outputs.get("BuildingCoverBaseRate"))
    if outputs.get("BuildingCoverUserRate"):
        summary["F21"] = _to_number(outputs.get("BuildingCoverUserRate"))


def build_excel_from_log(
    log_path: str | Path,
    output_path: str | Path,
    *,
    ruleset: str | None = None,
    template: str | Path | None = None,
) -> Path:
    """Build Policywise-style Excel with Building (U), BPP (V), Liability (W) factors.

    ``ruleset`` is accepted for CLI compatibility; when omitted/Building default,
    all three rating sections are filled.
    """
    _ = ruleset  # multi-section export is the default Policywise layout
    sections = collect_all_section_values(log_path)

    from logs.ruleset import extract_rulesets

    building_extract = extract_rulesets(str(log_path), ruleset_name="Building")
    header = {
        "policy_no": building_extract.header.policy_no,
        "module_id": building_extract.header.module_id,
        "project_id": building_extract.header.project_id,
    }
    inputs = {
        item.name: item.value
        for item in (building_extract.rulesets[0].inputs if building_extract.rulesets else [])
    }
    policy_no = header.get("policy_no") or inputs.get("PolNo") or Path(log_path).stem.split("_")[0]
    irpm = (sections.get("Building") or {}).get("irpm")

    workbook = load_workbook(resolve_template_path(template))
    fill_policy_header(workbook, policy_no=str(policy_no) if policy_no else None, irpm=irpm)
    fill_limits_from_building(workbook, log_path)

    for spec in SECTION_SPECS:
        fill_section_factors(workbook, spec, sections.get(spec.ruleset_name) or {})

    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(dest)
    return dest


def build_excel_from_payload(
    payload: dict[str, Any],
    output_path: str | Path,
    *,
    full_payload: dict[str, Any] | None = None,
    template: str | Path | None = None,
    log_path: str | Path | None = None,
) -> Path:
    """Write Excel from payloads; if log_path is given, fill all three sections."""
    if log_path:
        return build_excel_from_log(log_path, output_path, template=template)

    # Building-only fallback for callers that only have Building extract data.
    fields = {
        str(k): (None if v is None else str(v))
        for k, v in ((payload.get("rulesets") or [{}])[0].get("fields") or {}).items()
    }
    source = full_payload or payload
    header = payload.get("header") or {}
    inputs = {
        str(item.get("name")): str(item.get("value") or "")
        for item in (source.get("rulesets") or [{}])[0].get("inputs") or []
    }
    policy_no = header.get("policy_no") or inputs.get("PolNo")

    building_spec = SECTION_SPECS[0]
    values: dict[str, str | None] = {}
    for key, aliases in building_spec.field_aliases.items():
        values[key] = _lookup_aliases(fields, aliases)

    workbook = load_workbook(resolve_template_path(template))
    fill_policy_header(
        workbook,
        policy_no=str(policy_no) if policy_no else None,
        irpm=values.get("irpm"),
    )
    fill_section_factors(workbook, building_spec, values)

    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(dest)
    return dest


def create_blank_factor_workbook(output_path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "RatingFactors"
    ws["A1"] = "Section"
    ws["B1"] = "Factor"
    ws["C1"] = "Value"
    row = 2
    for name in RATING_SECTION_NAMES:
        ws.cell(row, 1, name)
        row += 1
    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
