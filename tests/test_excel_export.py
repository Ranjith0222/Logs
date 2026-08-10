from pathlib import Path

from openpyxl import load_workbook

from logs.coverages import (
    extract_coverage_rows,
    extract_rcflag_coverages,
    load_coverage_descriptions,
    resolve_liability_rating_limit,
)
from logs.excel_export import (
    build_excel_from_log,
    collect_all_section_values,
    collect_section_limits,
)

FIXTURE = Path(__file__).parent / "fixtures" / "building_ruleset_snippet.log"
COVERAGE_FIXTURE = Path(__file__).parent / "fixtures" / "coverage_rating_snippet.log"
FULL_SAMPLE = Path(__file__).parents[1] / "samples" / "uw_item_rating_building.log"
CLASS_CODE_TEMPLATE = Path("templates/class_code_template.xlsx")
LEGACY_TEMPLATE = Path("templates/policy_rating_template.xlsx")
BOP_COVERAGES = Path("templates/BOP_coverages.xlsx")


def test_build_excel_uses_uploaded_template(tmp_path: Path) -> None:
    template = LEGACY_TEMPLATE
    if not template.exists():
        template = Path("src/logs/data/policy_rating_template.xlsx")
    assert template.exists()
    out = tmp_path / "from_upload.xlsx"
    build_excel_from_log(FIXTURE, out, template=template)
    wb = load_workbook(out)
    assert "Build" in wb.sheetnames
    assert wb.active["U6"].value == 3.302
    # Building + BPP limits from Building ruleset inputs
    assert wb.active["F3"].value == 1516320
    assert wb.active["G3"].value == 5000
    assert wb.active["U3"].value == 1516320
    assert wb.active["V3"].value == 5000


def test_collect_section_limits_from_full_log() -> None:
    if not FULL_SAMPLE.exists():
        return
    limits = collect_section_limits(FULL_SAMPLE)
    assert limits["building"] == "1516320.00"
    assert limits["bpp"] == "5000.00"
    # SALES + BusinessCatgeory 18 → BuildingLimit (LiabilityPremiumNB1)
    assert limits["liability"] == "1516320.00"


def test_collect_bpp_and_liability_sections() -> None:
    if not FULL_SAMPLE.exists():
        return
    sections = collect_all_section_values(FULL_SAMPLE)
    bpp = sections["Business Personal Property"]
    liab = sections["Liability"]
    assert bpp["occ_rel"] == "3.257"
    assert bpp["construction"] == "0.825"
    assert bpp["loi"] == "1.767"
    assert bpp["deductible"] == "0.886"
    assert liab["liab_class"] == "2.974"
    assert liab["increased_limits"] == "1.074"
    assert liab["prop_damage_ded"] == "1"
    assert liab["base_lc"] == "0.008"


def test_build_excel_fills_bpp_and_liability_columns(tmp_path: Path) -> None:
    if not FULL_SAMPLE.exists():
        return
    template = LEGACY_TEMPLATE if LEGACY_TEMPLATE.exists() else None
    out = tmp_path / "multi.xlsx"
    build_excel_from_log(FULL_SAMPLE, out, template=template)
    wb = load_workbook(out)
    ws = wb[wb.sheetnames[0]]
    # Building U
    assert ws["U6"].value == 3.302
    # BPP V
    assert ws["V6"].value == 3.257
    assert ws["V7"].value == 0.825
    assert ws["V8"].value == 1.767
    # Liability W
    assert ws["W5"].value == 0.008
    assert ws["W17"].value == 2.974
    assert ws["W18"].value == 1.074
    assert ws["W19"].value == 1
    # Limits row: Building / BPP / Liability (left F/G/H and right U/V/W)
    assert ws["F3"].value == 1516320
    assert ws["G3"].value == 5000
    assert ws["H3"].value == 1516320
    assert ws["U3"].value == 1516320
    assert ws["V3"].value == 5000
    assert ws["W3"].value == 1516320


def test_load_bop_coverage_descriptions() -> None:
    path = BOP_COVERAGES if BOP_COVERAGES.exists() else Path("src/logs/data/BOP_coverages.xlsx")
    assert path.exists()
    mapping = load_coverage_descriptions(path)
    assert mapping["400101"] == "Building"
    assert mapping["400102"] == "Business Personal Property"
    assert mapping["400127"].startswith("Liability")


def test_resolve_liability_rating_limit_branches() -> None:
    building = "3141815.00"
    bpp = "10000.00"
    turnover = "390000"
    payroll = "800000.00"
    base_in = {"BuildingLimit": building, "BusiPersonalPropLimit": bpp}
    base_ev = {"AnnualTurnover": turnover, "Payroll": payroll}

    # SALES + category 18 → BuildingLimit
    assert (
        resolve_liability_rating_limit(
            {**base_in, "BusinessCatgeory": "18"},
            {},
            {**base_ev, "LiabilityExpBase": "SALES", "18InsuredStatus": "L"},
        )
        == building
    )
    # SALES + not 18 → AnnualTurnover
    assert (
        resolve_liability_rating_limit(
            {**base_in, "BusinessCatgeory": "14"},
            {},
            {**base_ev, "LiabilityExpBase": "SALES", "N18InsuredStatus": "O"},
        )
        == turnover
    )
    # LOI + occupant + BPP → BPP
    assert (
        resolve_liability_rating_limit(
            {**base_in, "BusinessCatgeory": "14"},
            {},
            {**base_ev, "LiabilityExpBase": "LOI", "N18InsuredStatus": "O"},
        )
        == bpp
    )
    # PAY + not 18 → Payroll
    assert (
        resolve_liability_rating_limit(
            {**base_in, "BusinessCatgeory": "14"},
            {},
            {**base_ev, "LiabilityExpBase": "PAY", "N18InsuredStatus": "O"},
        )
        == payroll
    )


def test_extract_coverage_rows_from_fixture() -> None:
    assert COVERAGE_FIXTURE.exists()
    rows = extract_coverage_rows(COVERAGE_FIXTURE, coverages_path=BOP_COVERAGES)
    by_code = {row.code: row for row in rows}
    assert by_code["400101"].sum_insured == "3141815.00"
    assert by_code["400101"].base_rate == "0.2110000"
    assert by_code["400101"].final_rate == "0.1780000"
    assert by_code["400101"].premium == "5593.0000000"
    assert by_code["400102"].sum_insured == "10000.00"
    assert by_code["400102"].premium == "115.0000000"
    # LiabilityPremiumNB1: SALES + category 18 → BuildingLimit
    assert by_code["400127"].sum_insured == "3141815.00"
    assert by_code["400127"].premium == "1100.0000000"
    assert by_code["400127"].label == "Liability"
    # Selected weSure Standard Endorsement Bundle (RCFlagN Y)
    assert by_code["400503"].label == "weSure Standard Endorsement Bundle"
    assert by_code["400503"].premium == "542.00"
    # Unselected covers must not appear even if residual premium exists
    assert "400136" not in by_code  # Hired Auto RCFlag=N
    assert "400440" not in by_code  # Non-owned Auto RCFlag=N
    assert {row.code for row in rows} == {"400101", "400102", "400127", "400503"}


def test_rcflagn_selected_flags() -> None:
    assert COVERAGE_FIXTURE.exists()
    rows = extract_rcflag_coverages(COVERAGE_FIXTURE)
    assert rows is not None
    by_code = {row.code: row for row in rows}
    assert by_code["400101"].selected is True
    assert by_code["400503"].selected is True
    assert by_code["400136"].selected is False
    assert by_code["400502"].selected is False  # Essential bundle not selected
    assert by_code["400503"].premium == "542.00"


def test_build_excel_fills_coverage_wise_table(tmp_path: Path) -> None:
    template = CLASS_CODE_TEMPLATE
    if not template.exists():
        template = Path("src/logs/data/class_code_template.xlsx")
    assert template.exists()
    assert COVERAGE_FIXTURE.exists()
    out = tmp_path / "coverage_wise.xlsx"
    build_excel_from_log(
        COVERAGE_FIXTURE,
        out,
        template=template,
        coverages_path=BOP_COVERAGES if BOP_COVERAGES.exists() else None,
    )
    ws = load_workbook(out).active
    # Refined template uses F/G/H for section factors.
    assert ws["K3"].value == "Coverage Code"
    assert ws["F3"].value == 3141815
    assert ws["G3"].value == 10000
    assert ws["H3"].value == 3141815
    # Coverage-wise table starts at row 4 under K-O.
    assert ws["K4"].value == "Building"
    assert ws["L4"].value == 3141815
    assert ws["M4"].value == 0.211
    assert ws["N4"].value == 0.178
    assert ws["O4"].value == 5593
    labels = [ws.cell(r, 11).value for r in range(4, 30) if ws.cell(r, 11).value]
    assert "Business Personal Property" in labels
    assert "Liability" in labels
    assert "weSure Standard Endorsement Bundle" in labels
    assert "Hired Auto Liability" not in labels
    assert "Business Income And Extra Expense – Revised Period Of Indemnity (in months)" not in labels
    liab_row = next(r for r in range(4, 30) if ws.cell(r, 11).value == "Liability")
    assert ws.cell(liab_row, 12).value == 3141815
    assert ws.cell(liab_row, 15).value == 1100
    bundle_row = next(
        r for r in range(4, 30) if ws.cell(r, 11).value == "weSure Standard Endorsement Bundle"
    )
    assert ws.cell(bundle_row, 15).value == 542
    # Total = 5593+115+1100+542
    assert "Total" in [ws.cell(r, 14).value for r in range(4, 30)]
    total_row = next(r for r in range(4, 30) if ws.cell(r, 14).value == "Total")
    assert ws.cell(total_row, 15).value == 7350
